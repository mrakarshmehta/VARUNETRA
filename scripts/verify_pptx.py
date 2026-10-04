import pptx
from pptx.util import Inches

prs = pptx.Presentation('VARUNETRA_FINAL_SUBMISSION/PPT/VARUNETRA_SIH_Final_Presentation.pptx')
w = prs.slide_width.inches
h = prs.slide_height.inches
print(f'Dimensions: {w:.3f}" x {h:.3f}" (Aspect Ratio: {w/h:.4f})')
assert abs(w/h - 16/9) < 0.01, 'Must be 16:9'

print(f'Total Slides: {len(prs.slides)}')
assert len(prs.slides) == 10, 'Must have exactly 10 slides'

overflows = []
for idx, slide in enumerate(prs.slides, 1):
    pic_count = 0
    tf_count = 0
    for shape in slide.shapes:
        right = (shape.left + shape.width) / 914400.0
        bottom = (shape.top + shape.height) / 914400.0
        if right > w + 0.1 or bottom > h + 0.1:
            overflows.append((idx, shape.name, round(right, 2), round(bottom, 2)))
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            pic_count += 1
        if shape.has_text_frame:
            tf_count += 1
    print(f'Slide {idx:02d}: shapes={len(slide.shapes)} | text_frames={tf_count} | pictures={pic_count}')

if overflows:
    print('OVERFLOWS DETECTED:', overflows)
else:
    print('[OK] Zero boundary overflows! All elements fit cleanly within 16:9 canvas.')
